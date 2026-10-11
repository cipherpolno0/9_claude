#!/usr/bin/env python3
"""สร้างรายงาน Excel จากผลของ rls_matrix.py (บทที่ 31)

ใช้: python3 -I scripts/security/rls_report.py <rls_matrix.json> <out.xlsx>
แผ่นงาน: สรุป / อ่าน / เพิ่ม / แก้ / ลบ (ตาราง x บทบาท) / บัญชีทดสอบ 5 บัญชี / นโยบาย
"""
import json
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

SRC, OUT = sys.argv[1], sys.argv[2]
d = json.load(open(SRC))
unit_scoped = set(d["unit_scoped"])
roles = d["role_order"]
role_names = d["role_names"]
results = d["results"]

FILL = {
    "own": PatternFill("solid", fgColor="D9F2D9"),
    "all": PatternFill("solid", fgColor="DCE9F7"),
    "none": PatternFill("solid", fgColor="F2F2F2"),
    "fn": PatternFill("solid", fgColor="FFF4D6"),
    "bad": PatternFill("solid", fgColor="F8D7DA"),
    "rule": PatternFill("solid", fgColor="EDE3F5"),
}
HEAD = Font(bold=True)
OWNER_TABLES = {"notifications": "เฉพาะของตน", "profiles": "ของตน + ผู้ที่ตนอนุมัติบัญชี", "quiz_attempts": "เฉพาะของตน (ผู้จัดการคลังเห็นทั้งหมด)"}


def is_num(v):
    return isinstance(v, int)


def cell(table, op, r):
    """คืน (ข้อความ, ชนิดสี)"""
    vals = r.get(op, {})
    if not vals:
        return "ไม่มีข้อมูลทดสอบ", "rule"
    errs = [v for v in vals.values() if isinstance(v, str) and v not in ("ok",)]
    if op == "delete":
        if all(v == "denied" for v in vals.values()):
            return "ไม่มี (ปิดใช้งานแทน)", "none"
        if any(is_num(v) and v > 0 for v in vals.values()):
            return "ลบได้!", "bad"
        if all(is_num(v) and v == 0 for v in vals.values()):
            return "—", "none"
        return "ติดกติกา*", "rule"
    if op == "insert":
        if all(v == "denied" for v in vals.values()):
            return "ผ่านฟังก์ชันเท่านั้น", "fn"
        okb = {b for b, v in vals.items() if v == "ok"}
        if table not in unit_scoped:
            if okb:
                return "ได้", "all"
            if all(v == "rls" for v in vals.values()):
                return "—", "none"
            return "ติดกติกา*", "rule"
        if okb == {"A", "B", "C"}:
            return "ทุกเขต", "all"
        if okb == {"A"}:
            return "เขตตน", "own"
        if okb & {"B", "C"}:
            return "ข้ามสายได้!", "bad"
        if all(v == "rls" for v in vals.values()):
            return "—", "none"
        return "ติดกติกา*", "rule"
    # read / update
    if any(v == "denied" for v in vals.values()):
        return ("ผ่านฟังก์ชันเท่านั้น" if op == "update" else "ไม่มีสิทธิ์ตาราง"), "fn"
    if errs:
        return "ติดกติกา*", "rule"
    seen = {b for b, v in vals.items() if is_num(v) and v > 0}
    if table not in unit_scoped:
        if table in OWNER_TABLES and op == "read" and not seen:
            return OWNER_TABLES[table], "own"
        return ("ได้", "all") if seen else ("—", "none")
    if seen == {"A", "B", "C"}:
        return "ทุกเขต", "all"
    if seen == {"A"}:
        return "เขตตน", "own"
    if not seen:
        return "—", "none"
    if "A" not in seen and seen <= {"B", "C"}:
        return "ข้ามสายได้!", "bad"
    return "บางเขต", "bad"


wb = Workbook()
ws = wb.active
ws.title = "สรุป"
bad_cells = []
for who, res in list(results["roles"].items()) + list(results["accounts"].items()):
    for t, r in res.items():
        for op in ("read", "insert", "update", "delete"):
            text, kind = cell(t, op, r)
            if kind == "bad" and not (who in ("admin", "central_staff")):
                bad_cells.append((who, t, op, text))
lines = [
    ("รายงานตรวจสิทธิ์ข้อมูลทุกตาราง (Row Level Security) บทที่ 31", True),
    ("วิธีตรวจ: สร้างข้อมูลทดสอบของทุกตารางใน 3 เขต A = ตำบล ก1-1 มหานิกาย (อยู่ใต้อำเภอ ก1 ของผู้ทดสอบ) / B = ตำบลในจังหวัดอื่น / C = ตำบลต่างนิกาย", False),
    ("แล้วลอง อ่าน เพิ่ม แก้ ลบ ในนามผู้ใช้จำลองบทบาทละ 1 คน (บทบาทที่มีเขต = อำเภอ ก1 มหานิกาย, ส่วนกลาง = หน่วยส่วนกลาง) และบัญชีทดสอบจริง 5 บัญชี ทุกการเขียนย้อนกลับ", False),
    ("ความหมาย: เขตตน = เฉพาะ A (หน่วยตนและใต้สังกัด) / ทุกเขต = A B C / — = ไม่ได้ / ผ่านฟังก์ชันเท่านั้น = ตารางไม่เปิดให้เขียนตรง ต้องผ่านฟังก์ชันที่ตรวจสิทธิ์เอง", False),
    ("ได้ = ข้อมูลกลางที่ไม่ผูกเขต (เช่น ปีการศึกษา หมวด ทะเบียนวัสดุ ข่าว) / ติดกติกา* = ถูกกติกาของข้อมูลปฏิเสธก่อนถึงการตรวจสิทธิ์ (ตรวจแยกในชุดทดสอบของบทนั้น)", False),
    (f"จำนวนตาราง {len(d['tables'])} ตาราง เปิด RLS ครบ {sum(1 for v in d['rls'].values() if v)} ตาราง ผูกเขตปกครอง {len(unit_scoped)} ตาราง", False),
    (f"ผลรวม: พบการเห็นหรือเขียนข้ามสายของบทบาทที่มีขอบเขตเขต {len(bad_cells)} จุด (ผู้ดูแลระบบและส่วนกลางเห็นทุกเขตตามที่ออกแบบ)", True),
]
for i, (text, bold) in enumerate(lines, 1):
    ws.cell(i, 1, text).font = Font(bold=bold)
for i, (who, t, op, text) in enumerate(bad_cells, len(lines) + 2):
    ws.cell(i, 1, f"{who} {t} {op}: {text}")
ws.column_dimensions["A"].width = 140

OPS = [("read", "อ่าน"), ("insert", "เพิ่ม"), ("update", "แก้"), ("delete", "ลบ")]
for op, title in OPS:
    s = wb.create_sheet(title)
    s.cell(1, 1, "ตาราง").font = HEAD
    s.cell(1, 2, "คำอธิบาย").font = HEAD
    s.cell(1, 3, "ผูกเขต").font = HEAD
    for j, rk in enumerate(roles, 4):
        c = s.cell(1, j, f"{role_names.get(rk, rk)}\n({rk})")
        c.font = HEAD
        c.alignment = Alignment(wrap_text=True, vertical="top")
        s.column_dimensions[get_column_letter(j)].width = 16
    for i, t in enumerate(d["tables"], 2):
        s.cell(i, 1, t)
        s.cell(i, 2, (d["comments"].get(t) or "")[:120])
        s.cell(i, 3, "ใช่" if t in unit_scoped else "")
        for j, rk in enumerate(roles, 4):
            r = results["roles"].get(rk, {}).get(t)
            if not r:
                s.cell(i, j, "ไม่มีข้อมูลทดสอบ")
                continue
            text, kind = cell(t, op, r)
            c = s.cell(i, j, text)
            c.fill = FILL[kind]
    s.column_dimensions["A"].width = 26
    s.column_dimensions["B"].width = 50
    s.column_dimensions["C"].width = 8
    s.freeze_panes = "D2"

s = wb.create_sheet("บัญชีทดสอบ 5 บัญชี")
accounts = list(results["accounts"])
s.cell(1, 1, "ตาราง (เฉพาะที่ผูกเขต)").font = HEAD
for j, a in enumerate(accounts, 2):
    s.cell(1, j, a).font = HEAD
    s.column_dimensions[get_column_letter(j)].width = 30
row = 2
fails = 0
for t in d["tables"]:
    if t not in unit_scoped:
        continue
    s.cell(row, 1, t)
    for j, a in enumerate(accounts, 2):
        r = results["accounts"][a].get(t)
        if not r:
            continue
        cross = []
        for op in ("read", "update", "delete"):
            for b in ("B", "C"):
                v = r[op].get(b)
                if is_num(v) and v > 0:
                    cross.append(f"{op}:{b}")
        for b in ("B", "C"):
            if r["insert"].get(b) == "ok":
                cross.append(f"insert:{b}")
        own = r["read"].get("A")
        text = ("ไม่ผ่าน " + ", ".join(cross)) if cross else ("ผ่าน (เห็นเขตตน)" if is_num(own) and own > 0 else "ผ่าน")
        fails += bool(cross)
        c = s.cell(row, j, text)
        c.fill = FILL["bad" if cross else ("own" if is_num(own) and own > 0 else "none")]
    row += 1
s.cell(row + 1, 1, f"ผล: {'ไม่พบการเห็นหรือเขียนข้ามสาย' if fails == 0 else f'พบ {fails} จุด'}").font = HEAD
s.column_dimensions["A"].width = 28
s.freeze_panes = "B2"

s = wb.create_sheet("นโยบาย")
for j, h in enumerate(["ตาราง", "RLS", "สิทธิ์ authenticated", "สิทธิ์ anon", "คำสั่ง", "บทบาทฐานข้อมูล", "เงื่อนไขอ่าน/แก้ (using)", "เงื่อนไขเขียน (with check)"], 1):
    s.cell(1, j, h).font = HEAD
row = 2
for t in d["tables"]:
    pols = d["policies"].get(t) or [{"cmd": "(ไม่มี policy = ปิดทั้งหมด)", "roles": "", "using": "", "check": ""}]
    for p in pols:
        g = d["grants"].get(t, {})
        vals = [t, "เปิด" if d["rls"].get(t) else "ปิด!", ",".join(sorted(g.get("authenticated", []))), ",".join(sorted(g.get("anon", []))),
                p["cmd"], p["roles"], (p["using"] or "")[:3000], (p["check"] or "")[:3000]]
        for j, v in enumerate(vals, 1):
            s.cell(row, j, v).alignment = Alignment(wrap_text=j >= 7, vertical="top")
        row += 1
for j, w in enumerate([26, 6, 22, 12, 10, 22, 70, 70], 1):
    s.column_dimensions[get_column_letter(j)].width = w
s.freeze_panes = "B2"
wb.save(OUT)
print("bad", len(bad_cells), "account_fails", fails)
